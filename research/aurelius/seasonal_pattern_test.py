"""
User's request: look for real patterns across Vanguard's and Aurelius's
actual MT5 trades - do losing months coincide between the two systems
and repeat across years, is there a day-of-week effect, and does
trading WITH or AGAINST the day's own early directional move predict
win/loss. Explicit ask to flag if we're over-reaching (data-dredging a
handful of thin slices until something looks significant by chance).

Ground truth: real MT5 deals from both EAs' own backtests (same
account 1301959345, GOLD#, ~2023-01 to 2026-09) - not a re-simulation.
Real M5 OHLC (engine.load_m5()) used only to derive each day's own
opening-range direction for the last question.
"""
import sys
sys.path.insert(0, ".")
import numpy as np
import pandas as pd
import openpyxl
from datetime import datetime, timedelta
from collections import defaultdict

VANGUARD_XLSX = "/root/.claude/uploads/0bd2ac72-7526-55cb-84f6-d8ea842f8c5b/10ce4f5a-20260921_-_ReportTester-1301959345_-_Vanguard_-_Backtest_1.xlsx"
AURELIUS_XLSX = "/root/.claude/uploads/0bd2ac72-7526-55cb-84f6-d8ea842f8c5b/56c08ddf-20260910_-_ReportTester-1301959345_-_Aurelius_M5.xlsx"


def find_deals_start(ws):
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == "Deals":
            return r + 2
    return None


def load_trades(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    start = find_deals_start(ws)
    open_stack = []
    trades = []
    for r in range(start, ws.max_row + 1):
        time = ws.cell(row=r, column=1).value
        if time is None:
            continue
        dtype = ws.cell(row=r, column=4).value
        direction = ws.cell(row=r, column=5).value
        profit = ws.cell(row=r, column=11).value
        if dtype not in ("buy", "sell"):
            continue
        t = datetime.strptime(time, "%Y.%m.%d %H:%M:%S")
        if direction == "in":
            open_stack.append((t, 1 if dtype == "buy" else -1))
        elif direction == "out":
            if open_stack:
                ot, odir = open_stack.pop(0)
                trades.append((ot, t, odir, profit))
    return trades


vanguard = load_trades(VANGUARD_XLSX)
aurelius = load_trades(AURELIUS_XLSX)
print(f"Vanguard: {len(vanguard)} trades ({vanguard[0][0].date()} to {vanguard[-1][1].date()})")
print(f"Aurelius: {len(aurelius)} trades ({aurelius[0][0].date()} to {aurelius[-1][1].date()})")

print("\n" + "=" * 70)
print("PART 1: monthly pattern - all years combined (calendar month, Jan-Dec)")
print("=" * 70)

def by_calendar_month(trades):
    by_m = defaultdict(list)
    for ot, ct, d, p in trades:
        by_m[ot.month].append(p)
    return by_m

vg_m = by_calendar_month(vanguard)
au_m = by_calendar_month(aurelius)

MONTH_NAMES = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
print(f"{'month':6}{'VG n':>6}{'VG net':>10}{'VG win%':>9}   {'AU n':>6}{'AU net':>10}{'AU win%':>9}")
vg_net_by_m = {}
au_net_by_m = {}
for m in range(1, 13):
    vp = vg_m.get(m, [])
    ap = au_m.get(m, [])
    vnet = sum(vp); anet = sum(ap)
    vg_net_by_m[m] = vnet; au_net_by_m[m] = anet
    vwr = 100 * np.mean([x > 0 for x in vp]) if vp else float('nan')
    awr = 100 * np.mean([x > 0 for x in ap]) if ap else float('nan')
    flag = ""
    if vnet < 0 and anet < 0: flag = "  <- BOTH LOSE"
    elif vnet < 0 or anet < 0: flag = "  <- one loses"
    print(f"{MONTH_NAMES[m]:6}{len(vp):6d}{vnet:10.2f}{vwr:9.1f}   {len(ap):6d}{anet:10.2f}{awr:9.1f}{flag}")

# correlation between monthly nets
months = list(range(1, 13))
vg_vec = np.array([vg_net_by_m[m] for m in months])
au_vec = np.array([au_net_by_m[m] for m in months])
corr = np.corrcoef(vg_vec, au_vec)[0, 1]
print(f"\ncorrelation between Vanguard's and Aurelius's monthly net profit: {corr:.3f}")

print("\n" + "=" * 70)
print("PART 1b: year x month grid - is a weak month weak EVERY year, or one bad year?")
print("=" * 70)

def year_month_grid(trades, label):
    grid = defaultdict(lambda: defaultdict(float))
    counts = defaultdict(lambda: defaultdict(int))
    for ot, ct, d, p in trades:
        grid[ot.year][ot.month] += p
        counts[ot.year][ot.month] += 1
    years = sorted(grid.keys())
    print(f"\n{label} - net $ by year x month:")
    header = "      " + "".join(f"{MONTH_NAMES[m]:>9}" for m in range(1, 13))
    print(header)
    for y in years:
        row = f"{y}  "
        for m in range(1, 13):
            v = grid[y].get(m, None)
            row += f"{v:9.0f}" if v is not None else f"{'--':>9}"
        print(row)

year_month_grid(vanguard, "VANGUARD")
year_month_grid(aurelius, "AURELIUS")

print("\n" + "=" * 70)
print("PART 2: day-of-week pattern (entry day)")
print("=" * 70)
DOW = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

def by_dow(trades):
    by_d = defaultdict(list)
    for ot, ct, d, p in trades:
        by_d[ot.weekday()].append(p)
    return by_d

vg_d = by_dow(vanguard)
au_d = by_dow(aurelius)
print(f"{'day':6}{'VG n':>6}{'VG net':>10}{'VG win%':>9}   {'AU n':>6}{'AU net':>10}{'AU win%':>9}")
for wd in range(7):
    vp = vg_d.get(wd, [])
    ap = au_d.get(wd, [])
    if not vp and not ap: continue
    vwr = 100 * np.mean([x > 0 for x in vp]) if vp else float('nan')
    awr = 100 * np.mean([x > 0 for x in ap]) if ap else float('nan')
    print(f"{DOW[wd]:6}{len(vp):6d}{sum(vp):10.2f}{vwr:9.1f}   {len(ap):6d}{sum(ap):10.2f}{awr:9.1f}")

print("\n" + "=" * 70)
print("PART 3: trading WITH vs AGAINST the day's own opening-range direction")
print("=" * 70)
print("""
Definition: each calendar day's 'opening direction' = sign of
(price at 08:00 server time - day's open price), using the first
completed M5 bar at/after 08:00 as the reference (08:00 is a
reasonable proxy for the London session open on this broker's server
time - GOLD# trades ~23h/day so there's no clean single 'day open'
the way equities have). A trade is 'WITH' the day if its own direction
matches the day's 08:00-vs-open sign; 'AGAINST' if it doesn't.
""")

import engine as E
df5 = E.load_m5()
df5["date"] = df5["time"].dt.date
day_open = df5.groupby("date")["open"].first()

# reference price: first close at/after 08:00 server time each day
df5["hour"] = df5["time"].dt.hour
ref = df5[df5["hour"] >= 8].groupby("date")["close"].first()

day_dir = {}
for d in day_open.index:
    if d in ref.index:
        day_dir[d] = 1 if ref[d] > day_open[d] else -1

def with_against(trades):
    withs, againsts, none_ = [], [], []
    for ot, ct, d, p in trades:
        dd = day_dir.get(ot.date())
        if dd is None:
            none_.append(p)
        elif dd == d:
            withs.append(p)
        else:
            againsts.append(p)
    return withs, againsts, none_

def report(label, trades):
    w, a, n_ = with_against(trades)
    print(f"\n{label}:")
    for name, ps in [("WITH day's 8am direction", w), ("AGAINST day's 8am direction", a)]:
        if not ps: continue
        wr = 100 * np.mean([x > 0 for x in ps])
        print(f"  {name:32s} n={len(ps):4d} net={sum(ps):10.2f} win%={wr:5.1f}")
    if n_:
        print(f"  (no reference available for {len(n_)} trades - skipped)")

report("VANGUARD", vanguard)
report("AURELIUS", aurelius)

print("\n" + "=" * 70)
print("PART 3b: confound check - is 'WITH' just a proxy for 'buy trades did better'?")
print("=" * 70)

def report_by_direction(label, trades):
    print(f"\n{label} - broken out by trade direction AND with/against:")
    for is_buy in (True, False):
        sub = [(ot, ct, d, p) for ot, ct, d, p in trades if (d > 0) == is_buy]
        w, a, n_ = with_against(sub)
        tag = "BUY" if is_buy else "SELL"
        for name, ps in [("WITH", w), ("AGAINST", a)]:
            if not ps: continue
            wr = 100 * np.mean([x > 0 for x in ps])
            print(f"  {tag:5s} {name:8s} n={len(ps):4d} net={sum(ps):10.2f} win%={wr:5.1f} avg={sum(ps)/len(ps):8.2f}")

report_by_direction("VANGUARD", vanguard)
report_by_direction("AURELIUS", aurelius)

print("\n" + "=" * 70)
print("PART 3c: significance check (proportion z-test on win rates, WITH vs AGAINST)")
print("=" * 70)
from math import sqrt

def prop_test(label, trades):
    w, a, _ = with_against(trades)
    n1, n2 = len(w), len(a)
    p1 = np.mean([x > 0 for x in w]); p2 = np.mean([x > 0 for x in a])
    p_pool = (sum(x > 0 for x in w) + sum(x > 0 for x in a)) / (n1 + n2)
    se = sqrt(p_pool * (1 - p_pool) * (1/n1 + 1/n2))
    z = (p1 - p2) / se if se > 0 else float('nan')
    print(f"{label}: win% WITH={100*p1:.1f} (n={n1})  win% AGAINST={100*p2:.1f} (n={n2})  z={z:.2f}  "
          f"({'|z|>1.96 -> real at 95% confidence' if abs(z) > 1.96 else 'not significant at 95%'})")

prop_test("VANGUARD", vanguard)
prop_test("AURELIUS", aurelius)

print("\n" + "=" * 70)
print("PART 3d: day-level permutation test (accounts for trades on the same day")
print("not being independent - a stricter, more honest test than the z-test above)")
print("=" * 70)

def with_against_net_gap(trades, day_dir_map):
    withs, againsts = [], []
    for ot, ct, d, p in trades:
        dd = day_dir_map.get(ot.date())
        if dd is None: continue
        (withs if dd == d else againsts).append(p)
    if not withs or not againsts: return 0.0
    return (sum(x > 0 for x in withs) / len(withs)) - (sum(x > 0 for x in againsts) / len(againsts))

def permutation_test(label, trades, n_perm=2000, seed=0):
    real_gap = with_against_net_gap(trades, day_dir)
    days = list(day_dir.keys())
    signs = np.array([day_dir[d] for d in days])
    rng = np.random.default_rng(seed)
    null_gaps = np.empty(n_perm)
    for i in range(n_perm):
        shuffled = rng.permutation(signs)
        perm_map = dict(zip(days, shuffled))
        null_gaps[i] = with_against_net_gap(trades, perm_map)
    pct = 100 * np.mean(null_gaps < real_gap)
    print(f"{label}: real WITH-AGAINST win%-gap={100*real_gap:.1f}pp  "
          f"null mean={100*null_gaps.mean():.1f}pp  percentile={pct:.1f}")

permutation_test("VANGUARD", vanguard)
permutation_test("AURELIUS", aurelius)
