"""
Parse a real MT5 Strategy Tester .xlsx report (Deals table) into round-trip
trades - the same methodology every version of MSG_Trader_EA.mq5's header
uses: walk the Deals table in order, open a round-trip on an 'in' deal, and
close it once cumulative 'out' volume equals the 'in' volume.

Profit is reported by MT5 in the ACCOUNT currency (ZAR on account 382043238),
so each round-trip also carries a currency-neutral P/L in USD
(price move x volume x contract size 100), which is what the Python
simulator produces and what real-vs-sim comparisons should use.
"""
import datetime as dt
import openpyxl

CONTRACT = 100.0


def _f(x):
    if x is None or x == "":
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    return float(str(x).replace(" ", "").replace("\xa0", ""))


def load_deals(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[0] == "Deals") + 2
    deals = []
    for r in rows[start:]:
        if r[0] is None or r[3] in (None, "balance"):
            continue
        if r[4] not in ("in", "out"):
            continue
        deals.append(dict(
            time=dt.datetime.strptime(r[0], "%Y.%m.%d %H:%M:%S"),
            type=r[3], dir=r[4], vol=_f(r[5]), price=_f(r[6]),
            profit=_f(r[10]), comment=r[12] or ""))
    return deals


def load_settings(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    out = {}
    for r in ws.iter_rows(values_only=True):
        for c in r:
            if isinstance(c, str) and "=" in c and c.startswith("Inp"):
                k, v = c.split("=", 1)
                out[k] = v
    return out


def round_trips(deals):
    trips, cur = [], None
    for d in deals:
        if d["dir"] == "in":
            assert cur is None, "overlapping positions - not single-slot"
            cur = dict(entry_time=d["time"], side=1 if d["type"] == "buy" else -1,
                       entry=d["price"], vol=d["vol"], legs=[], out_vol=0.0,
                       profit_ccy=0.0)
        else:
            cur["legs"].append(d)
            cur["out_vol"] += d["vol"]
            cur["profit_ccy"] += d["profit"]
            if abs(cur["out_vol"] - cur["vol"]) < 1e-9:
                cur["exit_time"] = d["time"]
                cur["pnl_usd"] = sum((l["price"] - cur["entry"]) * cur["side"] * l["vol"] * CONTRACT
                                     for l in cur["legs"])
                cur["nlegs"] = len(cur["legs"])
                cur["last_comment"] = cur["legs"][-1]["comment"]
                trips.append(cur)
                cur = None
    return trips


def summary(pnls):
    gp = sum(p for p in pnls if p > 0)
    gl = -sum(p for p in pnls if p < 0)
    return dict(n=len(pnls), net=round(sum(pnls), 2),
                pf=round(gp / gl, 3) if gl > 0 else float("inf"),
                win=round(100.0 * sum(1 for p in pnls if p > 0) / max(1, len(pnls)), 1))


if __name__ == "__main__":
    import sys
    t = round_trips(load_deals(sys.argv[1]))
    print(summary([x["pnl_usd"] for x in t]), "ccy:", summary([x["profit_ccy"] for x in t]))
    for x in t[:10]:
        print(x["entry_time"], x["side"], x["entry"], x["nlegs"], round(x["pnl_usd"], 2), x["last_comment"])


def entry_orders(path):
    """Entry order tickets: (time, side, SL, TP, requested px) - requested px
    is reconstructed exactly as (TP + 2*SL)/3 since TP = px +/- 2*risk."""
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    a = next(i for i, r in enumerate(rows) if r[0] == "Orders") + 2
    b = next(i for i, r in enumerate(rows) if r[0] == "Deals")
    out = []
    for r in rows[a:b]:
        if r[0] and isinstance(r[12], str) and r[12].startswith("MSG_MSG") and r[7] and r[8]:
            sl, tp = _f(r[7]), _f(r[8])
            out.append(dict(time=dt.datetime.strptime(r[0], "%Y.%m.%d %H:%M:%S"), side=r[3],
                            sl=sl, tp=tp, px=(tp + 2 * sl) / 3.0))
    return out
